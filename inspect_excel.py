#!/usr/bin/env python3
"""
Excel Workbook Inspector
Extracts cell mappings, formulas, and structure from Stall Chart.xlsm
"""

import openpyxl
from openpyxl.utils import get_column_letter
import json

def inspect_workbook(filename):
    """Inspect the Excel workbook and extract all relevant information"""
    
    print(f"Opening workbook: {filename}")
    wb = openpyxl.load_workbook(filename, data_only=False, keep_vba=True)
    
    print(f"\nWorkbook sheets: {wb.sheetnames}")
    
    results = {
        'sheets': {},
        'named_ranges': {},
        'formulas': {},
        'inputs': {},
        'outputs': {}
    }
    
    # Extract named ranges
    print("\n=== NAMED RANGES ===")
    for name, value in wb.defined_names.items():
        print(f"{name}: {value}")
        results['named_ranges'][name] = str(value)
    
    # Inspect each sheet
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n{'='*60}")
        print(f"SHEET: {sheet_name}")
        print(f"{'='*60}")
        
        sheet_data = {
            'dimensions': str(ws.dimensions),
            'cells': {},
            'formulas': {},
            'merged_cells': [str(cell_range) for cell_range in ws.merged_cells.ranges],
        }
        
        # Scan all cells with content
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None or cell.data_type == 'f':
                    cell_ref = f"{get_column_letter(cell.column)}{cell.row}"
                    
                    cell_info = {
                        'value': cell.value,
                        'data_type': cell.data_type,
                        'number_format': cell.number_format,
                    }
                    
                    # Check if it's a formula
                    if cell.data_type == 'f':
                        formula = f"={cell.value}"
                        cell_info['formula'] = formula
                        sheet_data['formulas'][cell_ref] = formula
                        
                        # Print formulas for inspection
                        if sheet_name == 'STALL CHART':
                            print(f"{cell_ref}: {formula}")
                    
                    sheet_data['cells'][cell_ref] = cell_info
        
        results['sheets'][sheet_name] = sheet_data
        
        # Print summary for this sheet
        print(f"\nTotal cells with content: {len(sheet_data['cells'])}")
        print(f"Total formulas: {len(sheet_data['formulas'])}")
        print(f"Merged cell ranges: {len(sheet_data['merged_cells'])}")
    
    return results

def identify_input_output_cells(results):
    """Analyze the workbook to identify input and output cells"""
    
    print("\n" + "="*60)
    print("IDENTIFYING INPUT/OUTPUT CELLS")
    print("="*60)
    
    # Look for patterns in STALL CHART sheet
    if 'STALL CHART' in results['sheets']:
        sheet_data = results['sheets']['STALL CHART']
        
        # Cells with values but no formulas are likely inputs
        for cell_ref, cell_info in sheet_data['cells'].items():
            if cell_info['data_type'] != 'f' and cell_info['value'] is not None:
                # Check if it's a numeric value (likely an input parameter)
                if isinstance(cell_info['value'], (int, float)):
                    print(f"Potential INPUT: {cell_ref} = {cell_info['value']}")
        
        # Cells with formulas are likely calculations/outputs
        print(f"\nTotal formula cells (outputs/calculations): {len(sheet_data['formulas'])}")

def export_to_json(results, filename='excel_structure.json'):
    """Export results to JSON file"""
    with open(filename, 'w') as f:
        json.dump(results, f, indent=2, default=str)
    print(f"\nResults exported to {filename}")

def main():
    workbook_path = 'Stall Chart.xlsm'
    
    try:
        results = inspect_workbook(workbook_path)
        identify_input_output_cells(results)
        export_to_json(results)
        
        # Create a focused summary for STALL CHART sheet
        if 'STALL CHART' in results['sheets']:
            print("\n" + "="*60)
            print("STALL CHART SHEET - KEY FORMULAS")
            print("="*60)
            
            formulas = results['sheets']['STALL CHART']['formulas']
            for cell_ref in sorted(formulas.keys(), key=lambda x: (int(''.join(filter(str.isdigit, x))), x)):
                print(f"{cell_ref}: {formulas[cell_ref]}")
        
    except Exception as e:
        print(f"Error: {e}")
        import traceback
        traceback.print_exc()

if __name__ == '__main__':
    main()
